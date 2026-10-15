"""Crea la plantilla de estilos (docs/manual-usuario/referencia.docx) con la que pandoc arma el .docx del manual.

Parte de la plantilla por defecto de pandoc y le da tamaño Carta, márgenes, fuentes, colores, encabezados que
empiezan en página nueva, pie con número de página y tablas con bordes. Solo hace falta volver a correrla si se
quiere cambiar el aspecto. Uso: python3 docx-plantilla.py <plantilla-por-defecto.docx> <salida.docx>
Requiere python-docx (pip install python-docx).
"""

import io
import sys

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFilter

# Tokens del diseño de la app (apps/web/src/styles/index.css): tinta, azul primario y sus tonos de vidrio.
AZUL = RGBColor(0x1D, 0x3F, 0xD1)
OSCURO = RGBColor(0x10, 0x22, 0x3A)
GRIS = RGBColor(0x4D, 0x5F, 0x79)
FUENTE = "Calibri"

ANCHO_PX, ALTO_PX = 1275, 1650  # Carta a 150 ppp


def fondo(intensidad):
    """El fondo de la app (--page): base azulada y tres manchas de color; `intensidad` lo suaviza para imprimir."""
    y, x = np.mgrid[0:ALTO_PX, 0:ANCHO_PX].astype(np.float32)
    lienzo = np.zeros((ALTO_PX, ANCHO_PX, 3), np.float32) + np.array([0xE6, 0xEF, 0xF7], np.float32)
    manchas = [  # de abajo hacia arriba: durazno, lila, verde agua
        ((60, 46, 28, 100), (0xFF, 0xCF, 0xA9)),
        ((48, 36, 100, 18), (0xC1, 0xB3, 0xFB)),
        ((55, 38, 12, 0), (0x8F, 0xE3, 0xD3)),
    ]
    for (rx, ry, cx, cy), color in manchas:
        t = np.sqrt(((x - cx / 100 * ANCHO_PX) / (rx / 100 * ANCHO_PX)) ** 2 + ((y - cy / 100 * ALTO_PX) / (ry / 100 * ALTO_PX)) ** 2)
        alfa = np.clip(1 - t / 0.62, 0, 1)[..., None] * intensidad
        lienzo = lienzo * (1 - alfa) + np.array(color, np.float32) * alfa
    return Image.fromarray(lienzo.clip(0, 255).astype(np.uint8))


def logo(lado):
    """El ícono de la app (apps/web/public/favicon.svg) dibujado a mano, con suavizado."""
    k, base = 8, lado * 8
    escala = base / 64
    img = Image.new("RGBA", (base, base), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((0, 0, base - 1, base - 1), radius=16 * escala, fill=(0x1D, 0x3F, 0xD1, 255))
    ancho = int(4 * escala)
    puntos = [(18, 24), (46, 24), (43, 44), (21, 44), (18, 24)]
    puntos = [(px * escala, py * escala) for px, py in puntos]
    d.line(puntos, fill="white", width=ancho, joint="curve")
    d.arc((26 * escala, 18 * escala, 38 * escala, 30 * escala), 180, 360, fill="white", width=ancho)
    for px, py in ((26, 24), (38, 24)):
        r = ancho / 2
        d.ellipse((px * escala - r, py * escala - r, px * escala + r, py * escala + r), fill="white")
    return img.resize((lado, lado), Image.LANCZOS)


def tarjeta_de_vidrio(imagen, caja, radio=34):
    """Dibuja un panel de vidrio (blanco translúcido, borde claro y sombra) sobre el fondo."""
    x0, y0, x1, y1 = caja
    forma = Image.new("L", imagen.size, 0)
    ImageDraw.Draw(forma).rounded_rectangle(caja, radius=radio, fill=255)
    sombra = Image.new("L", imagen.size, 0)
    ImageDraw.Draw(sombra).rounded_rectangle((x0, y0 + 14, x1, y1 + 14), radius=radio, fill=70)
    sombra = sombra.filter(ImageFilter.GaussianBlur(22))
    # La sombra solo asoma por fuera del panel: debajo del vidrio no debe oscurecerlo.
    sombra = Image.composite(Image.new("L", imagen.size, 0), sombra, forma)
    color = Image.new("RGBA", imagen.size, (0x1F, 0x32, 0x5A, 0))
    color.putalpha(sombra)
    imagen.alpha_composite(color)
    panel = Image.new("RGBA", imagen.size, (0, 0, 0, 0))
    ImageDraw.Draw(panel).rounded_rectangle(caja, radius=radio, fill=(255, 255, 255, 150), outline=(255, 255, 255, 235), width=3)
    imagen.alpha_composite(panel)


def jpeg(imagen):
    salida = io.BytesIO()
    imagen.convert("RGB").save(salida, "JPEG", quality=90, optimize=True)
    salida.seek(0)
    return salida


fondo_pagina = fondo(0.5)
portada = fondo(1.0).convert("RGBA")
# El panel rodea el título, la frase y la fecha de la portada (su posición sale del espaciado del estilo Title).
tarjeta_de_vidrio(portada, (90, 405, ANCHO_PX - 90, 780))
portada.alpha_composite(logo(190), (130, 160))

doc = Document(sys.argv[1])


def fuente(style, size, bold=None, color=None, italic=None):
    style.font.name = FUENTE
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FUENTE)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def parrafo(style, before=0, after=6, line=1.15, keep_next=None, align=None):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep_next is not None:
        fmt.keep_with_next = keep_next
    if align is not None:
        fmt.alignment = align


# Página Carta con márgenes cómodos para imprimir.
section = doc.sections[0]
section.page_width, section.page_height = Cm(21.59), Cm(27.94)
section.left_margin = section.right_margin = Cm(2.2)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)

# python-docx no encuentra por nombre los estilos de pandoc ("Heading 1"): se buscan recorriendo la lista.
styles = {style.name: style for style in doc.styles}
for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
    fuente(styles[name], 10.5, color=OSCURO)
parrafo(styles["Normal"], after=4)
parrafo(styles["Body Text"], after=6)
parrafo(styles["First Paragraph"], after=6)
parrafo(styles["Compact"], after=2)

fuente(styles["Title"], 32, bold=True, color=AZUL)
parrafo(styles["Title"], before=150, after=12, align=WD_ALIGN_PARAGRAPH.LEFT)
fuente(styles["Subtitle"], 14, color=GRIS)
parrafo(styles["Subtitle"], after=18, align=WD_ALIGN_PARAGRAPH.LEFT)
fuente(styles["Author"], 11, color=GRIS)
parrafo(styles["Author"], after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
fuente(styles["Date"], 11, color=GRIS)
parrafo(styles["Date"], after=2, align=WD_ALIGN_PARAGRAPH.LEFT)

# Cada parte (nivel 1) empieza en página nueva.
fuente(styles["Heading 1"], 24, bold=True, color=AZUL)
parrafo(styles["Heading 1"], before=0, after=12, keep_next=True)
styles["Heading 1"].paragraph_format.page_break_before = True
fuente(styles["Heading 2"], 16, bold=True, color=AZUL)
parrafo(styles["Heading 2"], before=18, after=6, keep_next=True)
fuente(styles["Heading 3"], 12.5, bold=True, color=OSCURO)
parrafo(styles["Heading 3"], before=12, after=4, keep_next=True)
fuente(styles["Heading 4"], 11, bold=True, color=OSCURO)
parrafo(styles["Heading 4"], before=8, after=3, keep_next=True)
fuente(styles["TOC Heading"], 20, bold=True, color=AZUL)
parrafo(styles["TOC Heading"], before=0, after=10)

# Capturas centradas y con su pie.
fuente(styles["Image Caption"], 9, italic=True, color=GRIS)
parrafo(styles["Image Caption"], before=2, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
parrafo(styles["Figure"], before=4, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
parrafo(styles["Captioned Figure"], before=4, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
fuente(styles["Caption"], 9, italic=True, color=GRIS)
fuente(styles["Table Caption"], 9, italic=True, color=GRIS)

# Notas (citas en el texto): recuadro azul claro con barra a la izquierda.
bloque = styles["Block Text"]
fuente(bloque, 10.5, color=OSCURO)
parrafo(bloque, before=6, after=8)
bloque.paragraph_format.left_indent = Cm(0.4)
bloque.paragraph_format.right_indent = Cm(0.2)
ppr = bloque.element.get_or_add_pPr()
for old in ppr.findall(qn("w:pBdr")) + ppr.findall(qn("w:shd")):
    ppr.remove(old)
ppr.append(
    parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="1D3FD1"/></w:pBdr>'
    )
)
ppr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="E4ECFB"/>'))

# Tablas con bordes finos y encabezado sombreado.
tabla = styles["Table"].element
tpr = tabla.find(qn("w:tblPr"))
if tpr is None:
    tpr = OxmlElement("w:tblPr")
    tabla.append(tpr)
for tag in ("w:tblBorders", "w:tblCellMar"):
    for old in tpr.findall(qn(tag)):
        tpr.remove(old)
tpr.append(
    parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        + "".join(
            f'<w:{lado} w:val="single" w:sz="4" w:space="0" w:color="C9D3E6"/>'
            for lado in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders>"
    )
)
tpr.append(
    parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:top w:w="60" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        '<w:bottom w:w="60" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>'
    )
)
# Celdas de vidrio claro: blanco casi opaco, como las tarjetas de la app sobre su fondo.
for old in tabla.findall(qn("w:tcPr")):
    tabla.remove(old)
tabla.append(parse_xml(f'<w:tcPr {nsdecls("w")}><w:shd w:val="clear" w:color="auto" w:fill="F8FAFD"/></w:tcPr>'))
for old in tabla.findall(qn("w:tblStylePr")):
    tabla.remove(old)
tabla.append(
    parse_xml(
        f'<w:tblStylePr {nsdecls("w")} w:type="firstRow"><w:rPr><w:b/></w:rPr>'
        '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="DCE5FB"/></w:tcPr></w:tblStylePr>'
    )
)

def fondo_en_cabecera(cabecera, imagen):
    """Pone la imagen como fondo de página completa: anclada detrás del texto, desde la esquina de la hoja."""
    cabecera.is_linked_to_previous = False
    parrafo_cab = cabecera.paragraphs[0] if cabecera.paragraphs else cabecera.add_paragraph()
    for run in list(parrafo_cab.runs):
        run._element.getparent().remove(run._element)
    run = parrafo_cab.add_run()
    forma = run.add_picture(jpeg(imagen), width=Cm(21.59), height=Cm(27.94))
    inline = forma._inline
    ancla = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>'
    )
    ancla.insert(3, inline.find(qn("wp:extent")))
    ancla.append(inline.find(qn("wp:docPr")))
    ancla.append(inline.find(qn("wp:cNvGraphicFramePr")))
    ancla.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, ancla)


section.different_first_page_header_footer = True
fondo_en_cabecera(section.header, fondo_pagina)
fondo_en_cabecera(section.first_page_header, portada)
section.first_page_footer.is_linked_to_previous = False

# Pie de página: nombre del manual y número de página.
pie = section.footer
pie.is_linked_to_previous = False
p = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
for run in list(p.runs):
    run._element.getparent().remove(run._element)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def campo(par, instruccion):
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, "1"), ("end", None)):
        run = par.add_run()
        run.font.size = Pt(9)
        run.font.color.rgb = GRIS
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._element.append(fc)
        elif texto == instruccion:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instruccion} "
            run._element.append(it)
        else:
            run.text = texto


texto = p.add_run("Manual de usuario · Página ")
texto.font.size = Pt(9)
texto.font.color.rgb = GRIS
campo(p, "PAGE")

doc.save(sys.argv[2])
print("Plantilla creada:", sys.argv[2])
